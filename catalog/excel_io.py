"""Read Excel Tables (ListObjects), whole sheets, and CSV into import payloads."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any

from openpyxl.utils import range_boundaries


MAX_PREVIEW_ROWS = 5
MAX_IMPORT_ROWS = 5000
MAX_COLS = 80


def _looks_like_excel_date_format(number_format: str) -> bool:
    fmt = (number_format or "").strip().lower()
    if not fmt or fmt in {"general", "@", "0", "0.00"}:
        return False
    if "fa-ir" in fmt or "fa_ir" in fmt:
        return True
    return any(tok in fmt for tok in ("yy", "mm", "dd", "yyyy", "m/", "d/", "/m", "/d"))


def _cell_str(value: Any, *, number_format: str = "") -> str:
    """Normalize a worksheet cell to a transferable string.

    Date/datetime values and Excel serials with a date (incl. fa-IR) number format
    become Jalali ``YYYY/MM/DD`` (e.g. ``1405/06/01``) so the system stores/display
    شمسی — never raw میلادی years in the Excel grid.
    """
    from catalog.jalali_dates import format_jalali_slash, excel_serial_to_gregorian

    if value is None:
        return ""
    if isinstance(value, datetime):
        return format_jalali_slash(value.date())
    if isinstance(value, date):
        return format_jalali_slash(value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if _looks_like_excel_date_format(number_format):
            as_date = excel_serial_to_gregorian(value)
            if as_date is not None:
                return format_jalali_slash(as_date)
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()
    return str(value).strip()


def _normalize_headers(raw: list[Any], width: int) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for i in range(width):
        base = _cell_str(raw[i]) if i < len(raw) else ""
        if not base:
            base = f"ستون {i + 1}"
        base = base[:120]
        n = seen.get(base, 0)
        seen[base] = n + 1
        headers.append(base if n == 0 else f"{base} ({n + 1})")
    return headers


def _split_header_rows(matrix: list[list[Any]]) -> tuple[list[str], list[list[str]]]:
    if not matrix:
        return [], []
    width = min(MAX_COLS, max((len(r) for r in matrix), default=0))
    if width <= 0:
        return [], []
    headers = _normalize_headers(matrix[0], width)
    rows: list[list[str]] = []
    for raw in matrix[1 : MAX_IMPORT_ROWS + 1]:
        row = [_cell_str(raw[i]) if i < len(raw) else "" for i in range(width)]
        if any(row):
            rows.append(row)
    return headers, rows


def detect_csv_delimiter(text: str) -> str:
    """Pick ``;`` or ``,`` from the first non-empty line (semicolon preferred when tied)."""
    line = ""
    for raw in text.splitlines():
        if raw.strip():
            line = raw
            break
    if not line:
        return ","
    semi = line.count(";")
    comma = line.count(",")
    if semi > 0 and semi >= comma:
        return ";"
    return ","


# Excel "File Origin → Arabic (Windows)" maps to Windows code page 1256.
CSV_ENCODING_LABELS = {
    "utf-8-sig": "UTF-8",
    "utf-8": "UTF-8",
    "cp1256": "Arabic (Windows)",
}


def decode_csv_bytes(content: bytes) -> tuple[str, str]:
    """Decode CSV bytes to text.

    Order matches common Iranian exports:
    1) UTF-8 (with/without BOM)
    2) Windows-1256 — same as Excel «Arabic (Windows)»
    """
    if content.startswith(b"\xef\xbb\xbf"):
        return content.decode("utf-8-sig"), "utf-8-sig"
    try:
        return content.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        return content.decode("cp1256"), "cp1256"
    except UnicodeDecodeError:
        pass
    # Last resort so preview still opens something readable.
    return content.decode("utf-8", errors="replace"), "utf-8"


def _csv_matrix(content: bytes) -> tuple[list[list[str]], str, str]:
    text, encoding = decode_csv_bytes(content)
    delimiter = detect_csv_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    matrix = [list(r) for r in reader]
    return matrix, delimiter, encoding


def _csv_display_name(uploaded_file) -> str:
    table_name = (getattr(uploaded_file, "name", "") or "Table1").rsplit("/", 1)[-1]
    if table_name.lower().endswith(".csv"):
        table_name = table_name[:-4] or "Table1"
    return table_name[:200] or "Table1"


def _matrix_from_table(ws, table) -> tuple[list[str], list[list[str]]]:
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    width = min(MAX_COLS, max_col - min_col + 1)
    if width <= 0:
        return [], []
    header_count = int(getattr(table, "headerRowCount", None) or 1)
    header_count = max(1, min(header_count, max_row - min_row + 1))

    # Prefer structured table column names when present
    col_names = []
    try:
        col_names = [c.name for c in (table.tableColumns or [])]
    except Exception:  # noqa: BLE001
        col_names = []
    if col_names and len(col_names) >= width:
        headers = _normalize_headers(col_names[:width], width)
    else:
        header_cells = [
            ws.cell(row=min_row, column=min_col + i).value for i in range(width)
        ]
        headers = _normalize_headers(header_cells, width)

    data_start = min_row + header_count
    rows: list[list[str]] = []
    for r in range(data_start, max_row + 1):
        if len(rows) >= MAX_IMPORT_ROWS:
            break
        # Skip totals row if marked
        totals = int(getattr(table, "totalsRowCount", None) or 0)
        if totals and r > max_row - totals:
            continue
        row = []
        for i in range(width):
            cell = ws.cell(row=r, column=min_col + i)
            row.append(_cell_str(cell.value, number_format=str(cell.number_format or "")))
        if any(row):
            rows.append(row)
    return headers, rows


def _matrix_from_sheet(ws) -> tuple[list[str], list[list[str]]]:
    """Read the used region of a worksheet as header + data rows."""
    matrix: list[list[Any]] = []
    for i, row in enumerate(ws.iter_rows(max_col=MAX_COLS)):
        if i > MAX_IMPORT_ROWS:
            break
        cells = [
            _cell_str(cell.value, number_format=str(cell.number_format or ""))
            for cell in row
        ]
        matrix.append(cells)
    # Trim trailing empty rows
    while matrix and not any(matrix[-1]):
        matrix.pop()
    # Trim trailing empty columns
    if matrix:
        width = max((len(r) for r in matrix), default=0)
        while width > 0 and all(
            (len(r) < width or not str(r[width - 1]).strip()) for r in matrix
        ):
            width -= 1
        matrix = [r[:width] for r in matrix]
    return _split_header_rows(matrix)


def _table_id(sheet_name: str, display_name: str) -> str:
    return f"{sheet_name}::{display_name}"


def _sheet_id(sheet_name: str) -> str:
    return f"sheet::{sheet_name}"


def inspect_workbook(uploaded_file) -> dict[str, Any]:
    """Return both Excel Tables and whole sheets for preview UI.

    Shape::
        {
          "tables": [...],
          "sheets": [...],
          "file_kind": "csv" | "xlsx",
          "csv_delimiter": ";" | "," | None,
          "csv_encoding": "utf-8" | "cp1256" | None,
          "csv_encoding_label": "UTF-8" | "Arabic (Windows)" | None,
        }
    """
    name = (getattr(uploaded_file, "name", "") or "").lower()
    content = uploaded_file.read()
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)

    if name.endswith(".csv"):
        matrix, delimiter, encoding = _csv_matrix(content)
        headers, rows = _split_header_rows(matrix)
        display = _csv_display_name(uploaded_file)
        enc_label = CSV_ENCODING_LABELS.get(encoding, encoding)
        item = {
            "name": display,
            "sheet_name": "CSV",
            "table_id": _table_id("CSV", display),
            "sheet_id": _sheet_id("CSV"),
            "ref": "",
            "headers": headers,
            "row_count": len(rows),
            "column_count": len(headers),
            "preview_rows": rows[:MAX_PREVIEW_ROWS],
            "kind": "table",
            "csv_delimiter": delimiter,
            "csv_encoding": encoding,
        }
        sheet_item = {
            **item,
            "kind": "sheet",
            "name": "CSV",
            "table_id": "",
            "sheet_id": _sheet_id("CSV"),
        }
        return {
            "tables": [item],
            "sheets": [sheet_item],
            "file_kind": "csv",
            "csv_delimiter": delimiter,
            "csv_encoding": encoding,
            "csv_encoding_label": enc_label,
        }

    from openpyxl import load_workbook

    # Tables are unavailable in read_only mode.
    wb = load_workbook(io.BytesIO(content), data_only=True)
    tables_out: list[dict] = []
    sheets_out: list[dict] = []
    try:
        for ws in wb.worksheets:
            sheet_title = str(ws.title)[:200]
            sheet_headers, sheet_rows = _matrix_from_sheet(ws)
            sheets_out.append({
                "name": sheet_title,
                "sheet_name": sheet_title,
                "table_id": "",
                "sheet_id": _sheet_id(sheet_title),
                "ref": "",
                "headers": sheet_headers,
                "row_count": len(sheet_rows),
                "column_count": len(sheet_headers),
                "preview_rows": sheet_rows[:MAX_PREVIEW_ROWS],
                "kind": "sheet",
                "table_count": len(list(getattr(ws, "tables", None) or {})),
            })
            if not getattr(ws, "tables", None):
                continue
            for key in list(ws.tables.keys()):
                table = ws.tables[key]
                display = str(getattr(table, "displayName", None) or table.name or key)
                headers, rows = _matrix_from_table(ws, table)
                tables_out.append({
                    "name": display[:200],
                    "sheet_name": sheet_title,
                    "table_id": _table_id(sheet_title, display),
                    "sheet_id": _sheet_id(sheet_title),
                    "ref": str(table.ref or ""),
                    "headers": headers,
                    "row_count": len(rows),
                    "column_count": len(headers),
                    "preview_rows": rows[:MAX_PREVIEW_ROWS],
                    "kind": "table",
                })
    finally:
        wb.close()
    return {
        "tables": tables_out,
        "sheets": sheets_out,
        "file_kind": "xlsx",
        "csv_delimiter": None,
        "csv_encoding": None,
        "csv_encoding_label": None,
    }


def preview_workbook(uploaded_file) -> list[dict]:
    """Backward-compatible: return Excel Table previews only."""
    return inspect_workbook(uploaded_file)["tables"]


def read_table_data(
    uploaded_file,
    *,
    sheet_name: str,
    table_name: str,
) -> tuple[list[str], list[list[str]]]:
    """Return headers+rows for one Excel Table by sheet + display name."""
    name = (getattr(uploaded_file, "name", "") or "").lower()
    content = uploaded_file.read()
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)

    if name.endswith(".csv"):
        matrix, _delimiter, _encoding = _csv_matrix(content)
        return _split_header_rows(matrix)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True)
    try:
        target_ws = None
        for ws in wb.worksheets:
            if str(ws.title) == sheet_name:
                target_ws = ws
                break
        if target_ws is None:
            raise ValueError(f"شیت «{sheet_name}» در فایل یافت نشد.")
        if not getattr(target_ws, "tables", None):
            raise ValueError(f"در شیت «{sheet_name}» هیچ Table یافت نشد.")

        target = None
        for key in target_ws.tables:
            table = target_ws.tables[key]
            display = str(getattr(table, "displayName", None) or table.name or key)
            if display == table_name or str(key) == table_name:
                target = table
                break
        if target is None:
            raise ValueError(
                f"جدول «{table_name}» در شیت «{sheet_name}» یافت نشد."
            )
        return _matrix_from_table(target_ws, target)
    finally:
        wb.close()


def read_sheet_data(uploaded_file, sheet_name: str) -> tuple[list[str], list[list[str]]]:
    """Return headers+rows for an entire worksheet (or CSV body)."""
    name = (getattr(uploaded_file, "name", "") or "").lower()
    content = uploaded_file.read()
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)
    if name.endswith(".csv"):
        matrix, _delimiter, _encoding = _csv_matrix(content)
        return _split_header_rows(matrix)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True)
    try:
        for ws in wb.worksheets:
            if str(ws.title) == sheet_name:
                return _matrix_from_sheet(ws)
        raise ValueError(f"شیت «{sheet_name}» در فایل یافت نشد.")
    finally:
        wb.close()
