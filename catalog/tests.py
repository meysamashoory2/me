"""Tests for Excel Table (ListObject) import, grid save, and report sources."""

from __future__ import annotations

import io
import json

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from catalog.models import ExcelTable, ExcelUpload
from reports.columns import get_column_groups, run_report

User = get_user_model()


def _make_xlsx_with_tables() -> bytes:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["کد", "نام", "موجودی"])
    ws.append(["A1", "قطعه یک", 10])
    ws.append(["A2", "قطعه دو", 5])
    ws["E1"] = "کد"
    ws["F1"] = "قیمت"
    ws["E2"] = "A1"
    ws["F2"] = 1000
    t1 = Table(displayName="Inventory", ref="A1:C3")
    t1.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t1)
    ws.add_table(Table(displayName="Prices", ref="E1:F2"))

    ws2 = wb.create_sheet("Extra")
    ws2.append(["X", "Y"])
    ws2.append([1, 2])
    # sheet without formal Table — must NOT appear in preview

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Kept for older helpers that only need a binary workbook
def _make_xlsx_bytes() -> bytes:
    return _make_xlsx_with_tables()


class ExcelManagementTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo")
        cls.admin = User.objects.get(username="admin")
        cls.expert = User.objects.get(username="expert")
        cls.viewer = User.objects.get(username="viewer")

    def _xlsx(self, name="sample.xlsx"):
        return SimpleUploadedFile(
            name,
            _make_xlsx_with_tables(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_list_requires_login(self):
        resp = self.client.get(reverse("excel_list"))
        self.assertEqual(resp.status_code, 302)

    def test_viewer_can_list_but_not_import(self):
        self.client.login(username="viewer", password="erp12345")
        self.assertEqual(self.client.get(reverse("excel_list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("excel_import")).status_code, 403)

    def test_preview_lists_excel_tables_not_plain_sheets(self):
        self.client.login(username="expert", password="erp12345")
        preview = self.client.post(
            reverse("excel_preview"),
            {"file": self._xlsx("multi.xlsx")},
        )
        self.assertEqual(preview.status_code, 200)
        data = preview.json()
        self.assertTrue(data["ok"])
        names = [t["name"] for t in data["tables"]]
        self.assertEqual(names, ["Inventory", "Prices"])
        self.assertNotIn("Extra", names)
        self.assertNotIn("Data", names)
        # Sheets are listed separately (including ones without Tables)
        sheet_names = [s["name"] for s in data["sheets"]]
        self.assertEqual(sheet_names, ["Data", "Extra"])
        self.assertEqual(data["sheet_count"], 2)
        self.assertEqual(data["table_count"], 2)
        inv = data["tables"][0]
        self.assertEqual(inv["sheet_name"], "Data")
        self.assertEqual(inv["row_count"], 2)
        self.assertEqual(inv["column_count"], 3)

    def test_import_selected_tables_by_name(self):
        self.client.login(username="expert", password="erp12345")
        confirm = self.client.post(
            reverse("excel_import_confirm"),
            {
                "file": self._xlsx("multi.xlsx"),
                "title": "فایل تست",
                "selected_sheets": json.dumps([
                    {"sheet": "Data", "table": "Inventory", "name": "جدول موجودی", "kind": "table"},
                    {"sheet": "Data", "table": "Prices", "name": "قیمت‌ها", "kind": "table"},
                ]),
            },
        )
        self.assertEqual(confirm.status_code, 200)
        payload = confirm.json()
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["table_count"], 2)
        upload = ExcelUpload.objects.get(pk=payload["upload_id"])
        tables = list(upload.tables.order_by("order"))
        self.assertEqual([t.name for t in tables], ["جدول موجودی", "قیمت‌ها"])
        self.assertEqual(tables[0].headers[:3], ["کد", "نام", "موجودی"])
        self.assertEqual(tables[0].row_count, 2)
        self.assertEqual(tables[1].headers, ["کد", "قیمت"])
        self.assertEqual(tables[1].rows[0][1], "1000")

        detail = self.client.get(reverse("excel_detail", args=[upload.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "جدول موجودی")
        self.assertContains(detail, "excel-grid")
        self.assertContains(detail, "دوبار کلیک")
        self.assertContains(detail, "جزئیات خطای انتقال")
        self.assertContains(detail, "excel-transfer-errors-dialog")
        self.assertNotContains(detail, "افزودن ردیف")
        self.assertNotContains(detail, "افزودن ستون")

    def test_import_whole_sheet_without_table(self):
        self.client.login(username="expert", password="erp12345")
        confirm = self.client.post(
            reverse("excel_import_confirm"),
            {
                "file": self._xlsx("multi.xlsx"),
                "title": "شیت کامل",
                "selected_sheets": json.dumps([
                    {"sheet": "Extra", "table": "", "name": "شیت اضافه", "kind": "sheet"},
                ]),
            },
        )
        self.assertEqual(confirm.status_code, 200)
        payload = confirm.json()
        self.assertTrue(payload["ok"])
        upload = ExcelUpload.objects.get(pk=payload["upload_id"])
        table = upload.tables.get()
        self.assertEqual(table.name, "شیت اضافه")
        self.assertEqual(table.sheet_name, "Extra")
        self.assertEqual(table.headers, ["X", "Y"])
        self.assertEqual(table.rows, [["1", "2"]])

    def test_semicolon_csv_preview_and_import(self):
        self.client.login(username="expert", password="erp12345")
        raw = "کد;نام;قیمت\nA1;قطعه;1000\nA2;قطعه دو;2000\n".encode("utf-8-sig")
        csv_file = SimpleUploadedFile("parts.csv", raw, content_type="text/csv")
        preview = self.client.post(reverse("excel_preview"), {"file": csv_file})
        self.assertEqual(preview.status_code, 200)
        data = preview.json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["file_kind"], "csv")
        self.assertEqual(data["csv_delimiter"], ";")
        self.assertIn(data.get("csv_encoding"), ("utf-8", "utf-8-sig"))
        self.assertEqual(len(data["tables"]), 1)
        self.assertEqual(len(data["sheets"]), 1)
        self.assertEqual(data["tables"][0]["headers"], ["کد", "نام", "قیمت"])
        self.assertEqual(data["tables"][0]["row_count"], 2)
        self.assertEqual(data["tables"][0]["column_count"], 3)

        csv_file2 = SimpleUploadedFile("parts.csv", raw, content_type="text/csv")
        confirm = self.client.post(
            reverse("excel_import_confirm"),
            {
                "file": csv_file2,
                "title": "CSV سمیکالن",
                "selected_sheets": json.dumps([
                    {"sheet": "CSV", "table": "", "name": "قطعات", "kind": "sheet"},
                ]),
            },
        )
        self.assertEqual(confirm.status_code, 200)
        payload = confirm.json()
        self.assertTrue(payload["ok"])
        table = ExcelUpload.objects.get(pk=payload["upload_id"]).tables.get()
        self.assertEqual(table.headers, ["کد", "نام", "قیمت"])
        self.assertEqual(table.rows[0], ["A1", "قطعه", "1000"])

    def test_arabic_windows_cp1256_csv(self):
        """Excel File Origin → Arabic (Windows) = cp1256."""
        from catalog.excel_io import decode_csv_bytes

        # cp1256 uses Arabic letter forms (ي / ك), like Iranian ERP exports.
        raw = (
            "کد کالا;شرح کالا;واحد\r\n"
            "07505045;سه راه تبديل 45 پروتکت;عدد\r\n"
        ).encode("cp1256")
        text, enc = decode_csv_bytes(raw)
        self.assertEqual(enc, "cp1256")
        self.assertIn("کد کالا", text)
        self.assertIn("سه راه تبديل", text)

        self.client.login(username="expert", password="erp12345")
        preview = self.client.post(
            reverse("excel_preview"),
            {"file": SimpleUploadedFile("inv.csv", raw, content_type="text/csv")},
        )
        self.assertEqual(preview.status_code, 200)
        data = preview.json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["csv_delimiter"], ";")
        self.assertEqual(data["csv_encoding"], "cp1256")
        self.assertEqual(data["csv_encoding_label"], "Arabic (Windows)")
        self.assertEqual(data["tables"][0]["headers"][:3], ["کد کالا", "شرح کالا", "واحد"])
        self.assertEqual(data["tables"][0]["preview_rows"][0][1], "سه راه تبديل 45 پروتکت")

        confirm = self.client.post(
            reverse("excel_import_confirm"),
            {
                "file": SimpleUploadedFile("inv.csv", raw, content_type="text/csv"),
                "title": "موجودی cp1256",
                "selected_sheets": json.dumps([
                    {"sheet": "CSV", "table": "", "name": "موجودی", "kind": "sheet"},
                ]),
            },
        )
        self.assertEqual(confirm.status_code, 200)
        table = ExcelUpload.objects.get(pk=confirm.json()["upload_id"]).tables.get()
        self.assertEqual(table.headers[:3], ["کد کالا", "شرح کالا", "واحد"])
        self.assertEqual(table.rows[0][0], "07505045")
        self.assertEqual(table.rows[0][1], "سه راه تبديل 45 پروتکت")

    def test_import_page_has_table_and_sheets_actions(self):
        self.client.login(username="expert", password="erp12345")
        resp = self.client.get(reverse("excel_import"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "مشاهده TABLE")
        self.assertContains(resp, "مشاهده SHEETS")
        self.assertContains(resp, "excel-view-tables")
        self.assertContains(resp, "excel-view-sheets")

    def test_missing_table_does_not_import(self):
        self.client.login(username="expert", password="erp12345")
        confirm = self.client.post(
            reverse("excel_import_confirm"),
            {
                "file": self._xlsx("multi.xlsx"),
                "title": "بد",
                "selected_sheets": json.dumps([
                    {"sheet": "Data", "table": "NoSuchTable", "name": "X"},
                ]),
            },
        )
        self.assertEqual(confirm.status_code, 400)
        self.assertFalse(confirm.json()["ok"])
        self.assertFalse(ExcelUpload.objects.filter(title="بد").exists())

    def test_save_table_headers_name_and_layout(self):
        self.client.login(username="expert", password="erp12345")
        upload = ExcelUpload.objects.create(title="ت", uploaded_by=self.expert)
        table = ExcelTable.objects.create(
            upload=upload,
            name="شیت۱",
            sheet_name="Data",
            headers=["A", "B"],
            rows=[["1", "2"]],
        )
        resp = self.client.post(
            reverse("excel_table_save", args=[table.pk]),
            data=json.dumps({
                "name": "جدول ویرایش‌شده",
                "headers": ["A", "B", "C"],
                "rows": [["1", "2", "3"], ["4", "5", "6"]],
                "layout": {"colWidths": [100, 140, 80], "rowHeights": [30, 32, 28]},
            }),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.json()["ok"])
        table.refresh_from_db()
        self.assertEqual(table.name, "جدول ویرایش‌شده")
        self.assertEqual(table.column_count, 3)
        self.assertEqual(table.layout.get("colWidths"), [100, 140, 80])

    def test_delete_table_manager_only(self):
        upload = ExcelUpload.objects.create(title="حذف", uploaded_by=self.admin)
        table = ExcelTable.objects.create(
            upload=upload, name="T1", sheet_name="Data", headers=["X"], rows=[["1"]],
        )
        self.client.login(username="expert", password="erp12345")
        forbidden = self.client.post(reverse("excel_table_delete", args=[table.pk]))
        self.assertEqual(forbidden.status_code, 403)
        self.assertTrue(ExcelTable.objects.filter(pk=table.pk).exists())

        self.client.login(username="admin", password="erp12345")
        ok = self.client.post(reverse("excel_table_delete", args=[table.pk]))
        self.assertEqual(ok.status_code, 302)
        self.assertFalse(ExcelTable.objects.filter(pk=table.pk).exists())

    def test_excel_tables_do_not_appear_as_direct_report_sources(self):
        """Reports must use transferred destinations, not raw Excel tables."""
        upload = ExcelUpload.objects.create(title="منبع", uploaded_by=self.admin)
        table = ExcelTable.objects.create(
            upload=upload,
            name="جدول فروش",
            sheet_name="Data",
            headers=["کد", "مقدار"],
            rows=[["C1", "9"], ["C2", "3"]],
        )
        groups = get_column_groups()
        ids = [g["id"] for g in groups]
        self.assertNotIn(table.source_id, ids)
        self.assertIn("history", ids)
        self.assertNotIn("file", ids)

        headers, rows, _payloads, _deeper = run_report(
            table.source_id,
            [
                {"key": "col_0", "source": table.source_id, "level": 1, "label": "کد"},
                {"key": "col_1", "source": table.source_id, "level": 1, "label": "مقدار"},
            ],
        )
        self.assertEqual(headers, [])
        self.assertEqual(rows, [])
